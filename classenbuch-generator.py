# -*- coding: utf-8 -*-
"""
Дневник Генератор - Финална версия с COM/win32com PDF Конвертиране (от Quittungs-Generator)
и фикс за .exe файлове (noconsole)
"""

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import pandas as pd
import os
import sys
import threading
import re
from datetime import datetime

# --- FIX ЗА .EXE ФАЙЛОВЕ (NOCONSOLE) ---
class DummyOutput:
    def write(self, x): pass
    def flush(self): pass

if sys.stdout is None:
    sys.stdout = DummyOutput()
if sys.stderr is None:
    sys.stderr = DummyOutput()
# --------------------------------------

# --- АВТОМАТИЧНО ЗАДАВАНЕ НА CURRENT DIRECTORY ---
def get_base_path():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.dirname(os.path.abspath(__file__))

APP_BASE_DIR = get_base_path()
os.chdir(APP_BASE_DIR)

import openpyxl
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# --- ИМПОРТИ ЗА PDF (С win32com както в Quittungs-Generator) ---
try:
    import win32com.client
    import pythoncom
    from pypdf import PdfWriter, PdfReader
    HAS_PDF_TOOLS = True
except ImportError:
    try:
        import win32com.client
        import pythoncom
        from PyPDF2 import PdfFileWriter as PdfWriter, PdfFileReader as PdfReader
        HAS_PDF_TOOLS = True
    except ImportError:
        HAS_PDF_TOOLS = False

from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

try:
    from PIL import Image, ImageTk, ImageDraw
    HAS_PILLOW = True
except ImportError:
    HAS_PILLOW = False

cancel_flag = False

def initialize_paths():
    try:
        default_files = {
            "schuelerliste": "schuelerliste.xlsx",
            "logo": "logo.png",
            "tool_logo": "tool_logo.png"
        }

        path_schueler = os.path.join(APP_BASE_DIR, default_files["schuelerliste"])
        if os.path.exists(path_schueler):
            excel_path_var.set(path_schueler)
            
        path_logo = os.path.join(APP_BASE_DIR, default_files["logo"])
        if os.path.exists(path_logo):
            logo_path_var.set(path_logo)

        path_tool_logo = os.path.join(APP_BASE_DIR, default_files["tool_logo"])
        
        if not os.path.exists(path_tool_logo):
            try:
                width, height = 500, 500
                img = Image.new("RGBA", (width, height), (255, 255, 255, 0))
                draw = ImageDraw.Draw(img)

                bg_white = "#FFFFFF"
                bg_green = "#00966E"
                bg_red = "#D62612"
                dark_bg = "#1F2937"

                draw.rounded_rectangle([25, 25, width - 25, height - 25], radius=100, fill=dark_bg, outline=bg_red, width=6)
                draw.rounded_rectangle([125, 125, 375, 390], radius=26, fill=bg_white, outline=bg_green, width=4)
                draw.rounded_rectangle([125, 125, 375, 190], radius=26, fill=bg_green)
                draw.rectangle([125, 160, 375, 190], fill=bg_green)

                draw.ellipse([155, 105, 185, 135], fill=bg_red)
                draw.ellipse([155, 380, 185, 410], fill=bg_red)
                draw.ellipse([315, 105, 345, 135], fill=bg_red)
                draw.ellipse([315, 380, 345, 410], fill=bg_red)

                line_color = "#E5E7EB"
                draw.rounded_rectangle([170, 230, 330, 245], radius=6, fill=line_color)
                draw.ellipse([135, 227, 150, 242], fill=bg_white, outline="#9CA3AF", width=2) 
                draw.rounded_rectangle([170, 280, 340, 295], radius=6, fill=line_color)
                draw.ellipse([135, 277, 150, 292], fill=bg_green) 
                draw.rounded_rectangle([170, 330, 280, 345], radius=6, fill=line_color)
                draw.ellipse([135, 327, 150, 342], fill=bg_red) 

                draw.rounded_rectangle([210, 145, 290, 160], radius=4, fill=bg_white)
                img.save(path_tool_logo)
            except Exception as e:
                pass

        if os.path.exists(path_tool_logo):
            tool_logo_path_var.set(path_tool_logo)

        default_templates_dir = os.path.join(APP_BASE_DIR, "templates")
        templates_dir_var.set(default_templates_dir)

        out_dir = os.path.join(APP_BASE_DIR, "out_dnevnik")
        output_dir_var.set(out_dir)
            
    except Exception as e:
        print(f"Грешка при инициализация: {e}")

def toggle_buttons(running=False):
    state = tk.DISABLED if running else tk.NORMAL
    cancel_state = tk.NORMAL if running else tk.DISABLED
    btn_generate.config(state=state)
    btn_pdf.config(state=state)
    btn_cancel.config(state=cancel_state)

def cancel_process():
    global cancel_flag
    cancel_flag = True
    status_var.set("Процесът се прекратява... Моля изчакайте.")
    btn_cancel.config(state=tk.DISABLED)

def show_message_threadsafe(title, msg, is_error=False):
    if is_error:
        messagebox.showerror(title, msg)
    else:
        messagebox.showinfo(title, msg)

def parse_filiale_and_class(val_str):
    s = str(val_str).strip()
    s_lower = s.lower()
    
    if any(w in s_lower for w in ['warteliste', 'abgemeldet', 'nan']):
        return "Неразпределени", s

    known_filialen = ['Остпарк', 'Весткройц']
    found_filiale = "Основен"
    class_part = s
    
    for f in known_filialen:
        if s.startswith(f):
            found_filiale = f
            class_part = s[len(f):].strip()
            break
            
    if not class_part:
        class_part = "Общ"
        
    return found_filiale, class_part

def get_template_filename(klasse_str):
    k_lower = str(klasse_str).lower().strip()
    
    # 1. Подготвителна група / Vorschule
    if any(w in k_lower for w in ['пуг', 'подготвителна', 'предшколна', 'vorschul', 'груп', 'подг']):
        return "dnevnik_potgotvitelna_grupa.docx"
        
    # 2. Извънкласни дейности / Клубове
    if any(w in k_lower for w in ['клуб', 'школа', 'танц', 'извънкласн']):
        return "dnevnik_izvanklasni_deinosti.docx"

    # 3. Числова проверка (1, 2, ..., 12)
    num_match = re.search(r'\d+', k_lower)
    if num_match:
        grade = int(num_match.group())
        if grade == 1:
            return "dnevnik_I_klas.docx"
        elif 2 <= grade <= 4:
            return "dnevnik_II_IV_klas.docx"
        elif 5 <= grade <= 12:
            return "dnevnik_V_XII_klas.docx"

    # 4. Текстова / Римска проверка за V - XII клас
    if any(w in k_lower for w in ['пети', 'шести', 'седми', 'осми', 'девети', 'десети', 'единадесети', 'дванадесети']):
        return "dnevnik_V_XII_klas.docx"
    if re.search(r'(?:^|[^a-z0-9])(xii|xi|x|ix|viii|vii|vi|v)(?=[^a-z0-9]|$)', k_lower):
        return "dnevnik_V_XII_klas.docx"

    # 5. Текстова / Римска проверка за II - IV клас
    if any(c in k_lower for c in ['втори', 'трети', 'четвърти']):
        return "dnevnik_II_IV_klas.docx"
    if re.search(r'(?:^|[^a-z0-9])(ii|iii|iv)(?=[^a-z0-9]|$)', k_lower):
        return "dnevnik_II_IV_klas.docx"

    # 6. Текстова / Римска проверка за I клас
    if 'първи' in k_lower:
        return "dnevnik_I_klas.docx"
    if re.search(r'(?:^|[^a-z0-9])i(?=[^a-z0-9]|$)', k_lower):
        return "dnevnik_I_klas.docx"

    return "dnevnik_V_XII_klas.docx"

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


# ==========================================
# 1. ГЕНЕРИРАНЕ НА WORD ДНЕВНИЦИ
# ==========================================
def start_generation():
    global cancel_flag
    excel_path = excel_path_var.get()
    templates_dir = templates_dir_var.get()
    output_dir = output_dir_var.get()
    
    if not excel_path or not templates_dir or not output_dir:
        messagebox.showerror("Грешка", "Моля изберете Schülerliste, папка с темплейти и изходна папка!")
        return

    cancel_flag = False
    toggle_buttons(running=True)
    progress_var.set(0)
    
    threading.Thread(target=process_diaries, args=(excel_path, templates_dir, output_dir), daemon=True).start()

def process_diaries(excel_path, templates_dir, output_dir):
    global cancel_flag
    try:
        excel_path = os.path.abspath(excel_path)
        templates_dir = os.path.abspath(templates_dir)
        output_dir = os.path.abspath(output_dir)

        root.after(0, lambda: status_var.set("Четене на ученическите списъци за Word..."))
        df = pd.read_excel(excel_path)
        
        name_col = 'Name Kind в кирилица' if 'Name Kind в кирилица' in df.columns else 'Name Kind in kyrillisch'
        if name_col not in df.columns:
            for col in df.columns:
                if 'kyrillisch' in str(col).lower() or 'кирилица' in str(col).lower():
                    name_col = col
                    break

        class_col = 'In Klasse' if 'In Klasse' in df.columns else None
        if not class_col:
            for col in df.columns:
                if 'klasse' in str(col).lower() or 'клас' in str(col).lower():
                    class_col = col
                    break

        df.dropna(subset=[class_col], inplace=True)
        
        parsed_filialen = []
        parsed_klassen = []
        for val in df[class_col]:
            f, k = parse_filiale_and_class(val)
            parsed_filialen.append(f)
            parsed_klassen.append(k)
            
        df['Parsed_Filiale'] = parsed_filialen
        df['Parsed_Klasse'] = parsed_klassen
        df = df[~df['Parsed_Klasse'].str.lower().isin(['warteliste', 'abgemeldet'])]

        if not os.path.exists(output_dir):
            os.makedirs(output_dir)

        grouped = df.groupby(['Parsed_Filiale', 'Parsed_Klasse'])
        total_groups = len(grouped)
        root.after(0, lambda: progress_bar.config(maximum=total_groups))

        idx = 0
        for (filiale, klasse), group_data in grouped:
            if cancel_flag: break
            
            msg = f"Генериране Word: {filiale}, клас {klasse}... ({idx+1}/{total_groups})"
            root.after(0, status_var.set, msg)
            
            filiale_dir = os.path.join(output_dir, str(filiale).replace("/", "_"))
            if not os.path.exists(filiale_dir):
                os.makedirs(filiale_dir)
                
            template_filename = get_template_filename(klasse)
            template_path = os.path.join(templates_dir, template_filename)

            if os.path.exists(template_path):
                doc = docx.Document(template_path)
            else:
                error_msg = f"ЛИПСВА ТЕМПЛЕЙТ!\n\nПрограмата се опитва да създаде дневник за клас '{klasse}'.\nЗа целта търси темплейт с име:\n'{template_filename}'\n\nПапката, в която го търси, е:\n{templates_dir}\n\nМоля, уверете се, че този файл съществува в папката!"
                raise FileNotFoundError(error_msg)

            group_data_sorted = group_data.sort_values(by=name_col, ascending=True, na_position='last')

            table_inserted = False
            for p in doc.paragraphs:
                if "{SPISAK_DECA}" in p.text:
                    p.text = ""  
                    table = doc.add_table(rows=1, cols=6)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    p._element.getparent().insert(p._element.getparent().index(p._element), table._element)
                    
                    header_tr = table.rows[0]._tr.get_or_add_trPr()
                    header_tr.append(OxmlElement('w:tblHeader'))
                    headers = ["№", "Име на детето (кирилица)", "Дата на раждане", "Място на раждане", "Адрес", "Родител / Контакт"]
                    col_widths = [Inches(0.5), Inches(2.2), Inches(1.3), Inches(1.5), Inches(2.0), Inches(2.5)]
                    
                    for i, title in enumerate(headers):
                        cell = table.rows[0].cells[i]
                        cell.width = col_widths[i]
                        set_cell_background(cell, "1A365D")
                        set_cell_margins(cell, top=120, bottom=120, left=80, right=80)
                        p_cell = cell.paragraphs[0]
                        p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p_cell.add_run(title)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
                        run.font.name = "Arial"
                    
                    for row_idx, (_, row) in enumerate(group_data_sorted.iterrows(), start=1):
                        row_cells = table.add_row().cells
                        trPr = table.rows[-1]._tr.get_or_add_trPr()
                        trPr.append(OxmlElement('w:cantSplit'))
                        row_data = [
                            str(row_idx),
                            str(row.get(name_col, row.get('Name Kind', ''))),
                            str(row.get('Geburtsdatum', '')).replace(' 00:00:00', ''),
                            str(row.get('Geburtsort', '')),
                            str(row.get('Eltern 1 - Wohnadresse', '')),
                            f"{row.get('Eltern 1 - Name', '')} ({row.get('Eltern 1 - Telefonnummer', '')})"
                        ]
                        bg_color = "F7FAFC" if row_idx % 2 == 0 else "FFFFFF"
                        for i, val in enumerate(row_data):
                            cell = row_cells[i]
                            cell.width = col_widths[i]
                            set_cell_background(cell, bg_color)
                            set_cell_margins(cell, top=90, bottom=90, left=80, right=80)
                            p_cell = cell.paragraphs[0]
                            p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                            run = p_cell.add_run(val)
                            run.font.size = Pt(9)
                            run.font.name = "Arial"
                            run.font.color.rgb = RGBColor(40, 40, 40)
                    
                    table_inserted = True
                    break

            if not table_inserted:
                doc.add_heading(f"Списък на учениците - Клас {klasse}", level=2)
                table = doc.add_table(rows=1, cols=6)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                header_tr = table.rows[0]._tr.get_or_add_trPr()
                header_tr.append(OxmlElement('w:tblHeader'))
                headers = ["№", "Име на детето (кирилица)", "Дата на раждане", "Място на раждане", "Адрес", "Родител / Контакт"]
                col_widths = [Inches(0.5), Inches(2.2), Inches(1.3), Inches(1.5), Inches(2.0), Inches(2.5)]
                
                for i, title in enumerate(headers):
                    cell = table.rows[0].cells[i]
                    cell.width = col_widths[i]
                    set_cell_background(cell, "1A365D")
                    set_cell_margins(cell, top=120, bottom=120, left=80, right=80)
                    p_cell = cell.paragraphs[0]
                    p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p_cell.add_run(title)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(9.5)
                    run.font.name = "Arial"
                
                for row_idx, (_, row) in enumerate(group_data_sorted.iterrows(), start=1):
                    row_cells = table.add_row().cells
                    trPr = table.rows[-1]._tr.get_or_add_trPr()
                    trPr.append(OxmlElement('w:cantSplit'))
                    row_data = [
                        str(row_idx),
                        str(row.get(name_col, row.get('Name Kind', ''))),
                        str(row.get('Geburtsdatum', '')).replace(' 00:00:00', ''),
                        str(row.get('Geburtsort', '')),
                        str(row.get('Eltern 1 - Wohnadresse', '')),
                        f"{row.get('Eltern 1 - Name', '')} ({row.get('Eltern 1 - Telefonnummer', '')})"
                    ]
                    bg_color = "F7FAFC" if row_idx % 2 == 0 else "FFFFFF"
                    for i, val in enumerate(row_data):
                        cell = row_cells[i]
                        cell.width = col_widths[i]
                        set_cell_background(cell, bg_color)
                        set_cell_margins(cell, top=90, bottom=90, left=80, right=80)
                        p_cell = cell.paragraphs[0]
                        p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                        run = p_cell.add_run(val)
                        run.font.size = Pt(9)
                        run.font.name = "Arial"
                        run.font.color.rgb = RGBColor(40, 40, 40)

            safe_klasse_name = str(klasse).replace("/", "_").replace("\\", "_")
            file_name = f"Дневник_{safe_klasse_name}_{filiale}.docx"
            file_path = os.path.join(filiale_dir, file_name)
            doc.save(file_path)
            
            idx += 1
            root.after(0, progress_var.set, idx)

        if cancel_flag:
            root.after(0, status_var.set, "Процесът е прекъснат.")
            return

        root.after(0, progress_var.set, total_groups)
        root.after(0, status_var.set, "Готово! Всички Word дневници са генерирани.")
        root.after(0, show_message_threadsafe, "Успех", f"Word дневниците бяха успешно генерирани в:\n{output_dir}")

    except Exception as e:
        root.after(0, status_var.set, "Грешка при Word генерация!")
        root.after(0, show_message_threadsafe, "Проблем с файл или темплейт", str(e), True)
    finally:
        root.after(0, toggle_buttons, False)


# ==========================================
# 2. КОНВЕРТИРАНЕ НА WORD В PDF И СБОРНИЦИ (Чрез win32com.client)
# ==========================================
def start_pdf_generation():
    global cancel_flag
    
    if not HAS_PDF_TOOLS:
        messagebox.showerror("Липсващи пакети", "Моля инсталирайте нужните пакети за PDF:\n\npip install pypdf pywin32")
        return

    output_dir = output_dir_var.get()
    
    if not output_dir or not os.path.exists(output_dir):
        messagebox.showerror("Грешка", "Моля първо изберете и генерирайте Word дневниците в изходната папка!")
        return

    cancel_flag = False
    toggle_buttons(running=True)
    progress_var.set(0)
    
    threading.Thread(target=process_pdf_conversion_and_merging, args=(output_dir,), daemon=True).start()

def process_pdf_conversion_and_merging(output_dir):
    global cancel_flag
    
    word_app = None
    # Инициализация на COM обектите за този тред
    pythoncom.CoInitialize()

    try:
        output_dir = os.path.abspath(output_dir)
        root.after(0, lambda: status_var.set("Подготовка за PDF конвертиране... Стартиране на Word във фонов режим."))
        
        # Настройка на шрифтове за ReportLab (кирилица)
        font_name = "Helvetica"
        font_bold_name = "Helvetica-Bold"
        try:
            arial_path = "C:\\Windows\\Fonts\\arial.ttf"
            arialbd_path = "C:\\Windows\\Fonts\\arialbd.ttf"
            if os.path.exists(arial_path) and os.path.exists(arialbd_path):
                pdfmetrics.registerFont(TTFont('CustomArial', arial_path))
                pdfmetrics.registerFont(TTFont('CustomArialBold', arialbd_path))
                font_name = "CustomArial"
                font_bold_name = "CustomArialBold"
        except Exception:
            pass

        styles = getSampleStyleSheet()
        separator_title_style = ParagraphStyle(
            'SeparatorTitle',
            parent=styles['Normal'],
            fontName=font_bold_name,
            fontSize=28,
            leading=34,
            alignment=1,
            textColor=colors.HexColor("#1A365D")
        )

        filialen = [d for d in os.listdir(output_dir) if os.path.isdir(os.path.join(output_dir, d))]
        
        if not filialen:
            root.after(0, show_message_threadsafe, "Грешка", "Няма намерени папки с филиали. Моля първо генерирайте Word дневниците!", True)
            root.after(0, toggle_buttons, False)
            return

        total_files = 0
        docx_map = {}
        for filiale in filialen:
            filiale_path = os.path.join(output_dir, filiale)
            docx_files = [f for f in os.listdir(filiale_path) if f.endswith('.docx') and not f.startswith('~')]
            docx_map[filiale] = docx_files
            total_files += len(docx_files)
            
        if total_files == 0:
            root.after(0, show_message_threadsafe, "Грешка", "Няма Word документи за конвертиране.", True)
            root.after(0, toggle_buttons, False)
            return

        root.after(0, lambda: progress_bar.config(maximum=total_files))

        # Стартираме скрит Word процес за конвертирането
        word_app = win32com.client.DispatchEx("Word.Application")
        word_app.Visible = False
        word_app.DisplayAlerts = 0 

        idx = 0
        for filiale in filialen:
            if cancel_flag: break
            
            filiale_path = os.path.join(output_dir, filiale)
            docx_files = docx_map[filiale]
            if not docx_files:
                continue

            writer = PdfWriter()

            for docx_file in docx_files:
                if cancel_flag: break
                
                parts = docx_file.replace("Дневник_", "").split("_")
                klasse_name = parts[0] if len(parts) > 0 else "Клас"

                msg = f"Конвертиране: {filiale} - {klasse_name}... ({idx+1}/{total_files})"
                root.after(0, status_var.set, msg)

                docx_path = os.path.join(filiale_path, docx_file)
                base_name = os.path.splitext(docx_file)[0]
                pdf_path = os.path.join(filiale_path, f"{base_name}.pdf")
                sep_pdf_path = os.path.join(filiale_path, f"sep_{base_name}.pdf")

                # 1. Създаване на разделителна страница в PDF
                doc_sep = SimpleDocTemplate(
                    sep_pdf_path,
                    pagesize=landscape(A4),
                    rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=30
                )
                story_sep = [
                    Spacer(1, 130),
                    Paragraph(f"Филиал: {filiale}", separator_title_style),
                    Spacer(1, 15),
                    Paragraph(f"Клас / Паралелка: {klasse_name}", ParagraphStyle('SubTitle', parent=separator_title_style, fontName=font_bold_name, fontSize=20, leading=24, textColor=colors.HexColor("#2C5282"))),
                    Spacer(1, 10),
                    Paragraph(f"Учебна година / Stand: {datetime.now().strftime('%d.%m.%Y')}", ParagraphStyle('Sub', parent=separator_title_style, fontName=font_name, fontSize=11, leading=14, textColor=colors.gray)),
                    PageBreak()
                ]
                doc_sep.build(story_sep)
                
                reader_sep = PdfReader(sep_pdf_path)
                for page in reader_sep.pages:
                    writer.add_page(page)

                # 2. СТАБИЛНО конвертиране през win32com
                try:
                    if not os.path.exists(pdf_path):
                        # Отваряме Word документа (ReadOnly=True)
                        doc = word_app.Documents.Open(docx_path, ReadOnly=True)
                        # Записваме като PDF (FileFormat=17 е wdFormatPDF)
                        doc.SaveAs(pdf_path, FileFormat=17)
                        # Затваряме без да записваме промените (0 = wdDoNotSaveChanges)
                        doc.Close(0)
                except Exception as e:
                    print(f"Грешка при конвертиране на {docx_file}: {e}")
                
                # Добавяме конвертирания PDF към сборника
                if os.path.exists(pdf_path):
                    reader_doc = PdfReader(pdf_path)
                    for page in reader_doc.pages:
                        writer.add_page(page)

                # Изтриваме помощната разделителна страница
                try:
                    if os.path.exists(sep_pdf_path):
                        os.remove(sep_pdf_path)
                except Exception:
                    pass
                
                idx += 1
                root.after(0, progress_var.set, idx)

            # 3. Записваме общия сборник за филиала
            if not cancel_flag:
                sbornik_path = os.path.join(output_dir, f"Дневници_Сборник_{filiale}.pdf")
                with open(sbornik_path, "wb") as output_file:
                    writer.write(output_file)

        if cancel_flag:
            root.after(0, status_var.set, "Процесът е прекъснат.")
            return

        root.after(0, progress_var.set, total_files)
        root.after(0, status_var.set, "Готово! Индивидуалните PDF и сборници са създадени.")
        root.after(0, show_message_threadsafe, "Успех", f"Всички индивидуални PDF дневници и общите сборници по филиали са успешно конвертирани и запазени в папка:\n{output_dir}")

    except Exception as e:
        root.after(0, status_var.set, "Грешка при PDF обработка!")
        root.after(0, show_message_threadsafe, "Критична грешка", f"Възникна грешка:\n{e}", True)
    finally:
        # Задължително затваряне на Word приложението и COM нишката
        if word_app is not None:
            try:
                word_app.Quit()
            except:
                pass
        pythoncom.CoUninitialize()
        root.after(0, toggle_buttons, False)


# ==========================================
# ГРАФИЧЕН ИНТЕРФЕЙС (GUI)
# ==========================================
def select_excel():
    current_dir = os.getcwd()
    path = filedialog.askopenfilename(filetypes=[("Excel файлове", "*.xlsx *.xls")])
    os.chdir(current_dir)
    if path:
        excel_path_var.set(path)

def select_templates_dir():
    current_dir = os.getcwd()
    path = filedialog.askdirectory()
    os.chdir(current_dir)
    if path:
        templates_dir_var.set(path)

def select_out_dir():
    current_dir = os.getcwd()
    path = filedialog.askdirectory()
    os.chdir(current_dir)
    if path:
        output_dir_var.set(path)

root = tk.Tk()
root.title("Дневник Генератор - Стабилна версия")
root.geometry("700x680")

excel_path_var = tk.StringVar()
templates_dir_var = tk.StringVar()
output_dir_var = tk.StringVar()
logo_path_var = tk.StringVar()
tool_logo_path_var = tk.StringVar()

frame = tk.Frame(root, padx=15, pady=15)
frame.pack(expand=True, fill=tk.BOTH)
frame.grid_columnconfigure(0, weight=1)

initialize_paths()

logo_frame = tk.Frame(frame)
logo_frame.grid(row=0, column=0, columnspan=2, pady=(0, 10))

try:
    resample_filter = Image.Resampling.LANCZOS
except AttributeError:
    try:
        resample_filter = Image.LANCZOS
    except AttributeError:
        resample_filter = Image.ANTIALIAS

school_logo = logo_path_var.get()
if os.path.exists(school_logo):
    try:
        img1 = Image.open(school_logo)
        img1.thumbnail((200, 60), resample_filter)
        tk_img1 = ImageTk.PhotoImage(img1)
        root.tk_img1 = tk_img1
        tk.Label(logo_frame, image=tk_img1).pack(side=tk.LEFT, padx=5)
    except Exception as e:
        pass

tool_logo = tool_logo_path_var.get()
if os.path.exists(tool_logo):
    try:
        img2 = Image.open(tool_logo)
        img2.thumbnail((60, 60), resample_filter)
        tk_img2 = ImageTk.PhotoImage(img2)
        root.tk_img2 = tk_img2
        tk.Label(logo_frame, image=tk_img2).pack(side=tk.LEFT, padx=5)
    except Exception as e:
        pass

tk.Label(frame, text="Дневник Генератор", font=("Helvetica", 16, "bold"), fg="#1A365D").grid(row=1, column=0, columnspan=2, pady=(5, 15))

tk.Label(frame, text="1. Изберете Schülerliste (Excel):", font=("Helvetica", 10, "bold")).grid(row=2, column=0, sticky="w", pady=2)
tk.Entry(frame, textvariable=excel_path_var, width=54).grid(row=3, column=0, padx=(0, 5), sticky="ew")
tk.Button(frame, text="Преглед...", command=select_excel).grid(row=3, column=1)

tk.Label(frame, text="2. Изберете папка с темплейти:", font=("Helvetica", 10, "bold")).grid(row=4, column=0, sticky="w", pady=(10, 2))
tk.Entry(frame, textvariable=templates_dir_var, width=54).grid(row=5, column=0, padx=(0, 5), sticky="ew")
tk.Button(frame, text="Преглед...", command=select_templates_dir).grid(row=5, column=1)

tk.Label(frame, text="3. Изберете Изходна папка:", font=("Helvetica", 10, "bold")).grid(row=6, column=0, sticky="w", pady=(10, 2))
tk.Entry(frame, textvariable=output_dir_var, width=54).grid(row=7, column=0, padx=(0, 5), sticky="ew")
tk.Button(frame, text="Преглед...", command=select_out_dir).grid(row=7, column=1)

btn_frame = tk.Frame(frame)
btn_frame.grid(row=8, column=0, columnspan=2, pady=(20, 10))

btn_generate = tk.Button(btn_frame, text="🚀 Генерирай Word дневници", font=("Helvetica", 10, "bold"), command=start_generation, bg="#4CAF50", fg="white")
btn_generate.pack(side=tk.LEFT, padx=5, ipadx=8, ipady=5)

btn_pdf = tk.Button(btn_frame, text="📄 Конвертирай & Обедини PDF", font=("Helvetica", 10, "bold"), command=start_pdf_generation, bg="#2B6CB0", fg="white")
btn_pdf.pack(side=tk.LEFT, padx=5, ipadx=8, ipady=5)

btn_cancel = tk.Button(btn_frame, text="🛑 Спри", font=("Helvetica", 10, "bold"), command=cancel_process, bg="#cc3025", fg="white", state=tk.DISABLED)
btn_cancel.pack(side=tk.LEFT, padx=5, ipadx=8, ipady=5)

status_var = tk.StringVar()
status_var.set("Готов за работа.")
tk.Label(frame, textvariable=status_var, fg="blue", font=("Helvetica", 10)).grid(row=9, column=0, columnspan=2, pady=(5, 4))

progress_var = tk.IntVar()
progress_bar = ttk.Progressbar(frame, variable=progress_var, mode='determinate')
progress_bar.grid(row=10, column=0, columnspan=2, sticky="ew", pady=(0, 10))

tk.Label(frame, text="© 2026 - Дневник Генератор (COM стабилна версия)", font=("Helvetica", 8), fg="gray").grid(row=11, column=0, columnspan=2, pady=(0, 2))

root.mainloop()